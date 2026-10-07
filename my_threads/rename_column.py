import pandas as pd
from openpyxl.styles import Alignment 
import os

from PySide6 import QtWidgets, QtCore
from colorama import Fore
from openpyxl import load_workbook
import openpyxl
import os
import pyperclip
from PySide6 import QtWidgets, QtCore
from colorama import Fore
import global_vars 
import os
import pyperclip
import pandas as pd
from datetime import datetime
from my_threads.functions import get_files_and_sheets_from_pyperclip, all_control_elements_on, all_control_elements_off, check_files_modified, check_excel_file_is_open, on_finsh_change_thread
from time import sleep


class RenameColumnThread(QtCore.QThread):
    def __init__ (self, md_files=False, parent=None):
        QtCore.QThread.__init__(self, parent)
        self.message_title = "Изменение заголовков в выбранных листах."
        self.md_files = md_files

    mysignal_info_label = QtCore.Signal(str)
    mysignal_finished = QtCore.Signal(str, str, str, str, list)


    def run(self):

        self.error_message = "Не удалось найти ни одного заголовка для изменения!"
        self.warning_message = ""
        self.info_message = ""
        self.md_files_opened = []


        self.is_src_files_modifyed = check_files_modified('.Исходники')


        if self.is_src_files_modifyed:
            self.error_message = ('В папку .Исходники были добавлены новые файлы или\n'
                                  'некоторые файлы в ней были пересохранены или удалены.\n'
                                  'Нажмите кнопку "Просмотерь разметку"!')

            # global_vars.ui.info_label.setText(self.error_message)
            return
        

        if global_vars.ui.lineEditOldColumnNameInHeader.text() == '' and global_vars.ui.lineEditNewColumnNameInHeader.text() == '':
            self.error_message = "Выберите что на что нужно поменять!"
            return
        if (
            global_vars.ui.lineEditOldColumnNameInHeader.text() == global_vars.ui.lineEditNewColumnNameInHeader.text() and
            global_vars.ui.radioButtonOldInTopHeader.isChecked() == global_vars.ui.radioButtonNewInTopHeader.isChecked()
            ):
            self.error_message = "Новое значение такое же как старое в той же строке заголовка!"
            return
        
        

        changed_qty = 0

        if global_vars.ui.radioButtonOldInTopHeader.isChecked():
            old_header_row_number = 1
        else:
            old_header_row_number = 2

        if global_vars.ui.radioButtonNewInTopHeader.isChecked():
            new_header_row_number = 1
        else:
            new_header_row_number = 2            

        files_sheets_list = get_files_and_sheets_from_pyperclip()
        files_list = list({files_sheets[0] for files_sheets in files_sheets_list})
        files_list.sort()

        # self.md_files_opened = [file for file in files_list if check_excel_file_is_open(file)]
        self.md_files_opened = []
        
        for file_number, file in enumerate(files_list):
            self.mysignal_info_label.emit(
                f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. {file_number} из {len(files_list)}. "
                f"Переименование заголовков. Проверяем не открыт ли на рабочем столе: {file}.",
                'color: blue')
            #sleep(0.01)
            if check_excel_file_is_open(file):
                self.md_files_opened.append(file)

        if self.md_files_opened:
            self.warning_message = (
                f"Не можем поменять заголовки в следующих файлах\n"
                f"так как они открыты на рабочем столе:\n"
                f"{'\n'.join(self.md_files_opened)}"
                )
            print('B')
            return

        file_preceding = ""
        need_to_save = False

        for file_sheet_number, file_sheet_list in enumerate(files_sheets_list, 1):
            # print(Fore.MAGENTA, file_sheet_list,  Fore.RESET)
            file = file_sheet_list[0]


            if file != file_preceding:

                if need_to_save:
                    self.mysignal_info_label.emit(
                        f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. { files_list.index(file)+1 } из {len(files_list)}. '
                        f'Сохраняем с измененными заголовками книгу "{file}"',
                        'color: blue')
                    # print(Fore.GREEN, 'Мы тут', Fore.RESET)
                    wb.save(os.path.join(global_vars.project_folder, '.Размеченные', file_preceding))

                file_preceding = file
                need_to_save = False

                self.mysignal_info_label.emit(
                        f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. { files_list.index(file)+1 } из {len(files_list)}. '
                        f'Загружаем для переименования заголовков книгу "{file}"',
                        'color: blue')
                wb = load_workbook(os.path.join(global_vars.project_folder, '.Размеченные', file))

            
            sheet_name = file_sheet_list[1]


            self.mysignal_info_label.emit(
                f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. '
                f'{file_sheet_number} из {len(files_sheets_list)}. '
                f'Сканируем заголовки в книге "{file}" в листе "{ sheet_name }"',
                'color_blue')
            #sleep(0.01)

            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                self.info_message = f"Изменено заголовков: {changed_qty}."   
                continue
                
            old_header_row = list(ws.iter_rows(values_only=True))[old_header_row_number-1]
            # print(Fore.GREEN, old_header_row, Fore.RESET)
   
            for col_number, old_header_cell in enumerate(old_header_row, 1):
                if col_number in (1,2): continue

                
                if (str(old_header_cell) == str(global_vars.ui.lineEditOldColumnNameInHeader.text()) or
                    old_header_cell == None and str(global_vars.ui.lineEditOldColumnNameInHeader.text()) == ''):
                    
                    # print(Fore.YELLOW, old_header_cell == global_vars.ui.lineEditOldColumnNameInHeader.text(), Fore.RESET)
                   
                    ws.cell(row=old_header_row_number, column=col_number, value='')                
                    ws.cell(row=new_header_row_number, column=col_number, value=global_vars.ui.lineEditNewColumnNameInHeader.text())
                    need_to_save = True
                    changed_qty += 1
                    self.error_message = ""


        if need_to_save:
            self.mysignal_info_label.emit(
                f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. { files_list.index(file)+1 } из {len(files_list)}. '
                f'Сохраняем с измененными заголовками книгу "{file}"',
                'color: blue')
            wb.save(os.path.join(global_vars.project_folder, '.Размеченные', file_preceding))
        else:
            wb.close()

        self.info_message = f"Изменено заголовков: {changed_qty}."     


        self.mysignal_finished.emit(
            self.message_title,
            self.error_message,
            self.warning_message,
            self.info_message,
            self.md_files_opened
        )


    def on_clicked(self):
     
        self.start() # Запускаем поток  
     

    #def on_finished(self): # Вызывается при завершении потока
    #    on_finsh_change_thread(self.message_title, self.error_message, self.warning_message, self.info_message, self.md_files_opened)
    #    all_control_elements_on()



