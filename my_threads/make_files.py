from PySide6 import QtWidgets, QtCore
from colorama import Fore
from datetime import datetime
import global_vars 
import os, random
import pandas as pd
from time import sleep
from my_threads.functions import check_files_modified, pop_up_files
from openpyxl import load_workbook, styles
from openpyxl.utils.cell import get_column_letter
from my_threads.functions import all_control_elements_off, all_control_elements_on, check_excel_file_is_open

class MakeFilesThread(QtCore.QThread):
 
    mysignal_info_label = QtCore.Signal(str, str)

    def __init__ (self, parent=None):
        QtCore.QThread.__init__(self, parent) 
        self.message_title = "Создаём файлы"
        self.error_message = ""
        self.warning_message = ""
        self.result_df_len = 0

        
    def clean_folder_marked(self):

        errors_list = []


        source_files = list(os.walk(os.path.join(global_vars.project_folder,'.Исходники')))[0][2]
        marked_files = [file for file in list(os.walk(os.path.join(global_vars.project_folder,'.Размеченные')))[0][2] if file[0] != '~']
     
        for file in marked_files:
            #sleep(0.0001)
            self.mysignal_info_label.emit(
                f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                f"Проверяем наличие файла {file} из ./Размеченные в .Исходники/",
                'color: blue')  
                       
            if file[3:] not in source_files: # file[3:] чтобы откусить приставку md_ в начале
                try:
                    os.remove(os.path.join(global_vars.project_folder, '.Размеченные', file))
                except PermissionError:
                    errors_list.append("Книга {file} есть в папке .Размеченные/,\
                            но ей нет соответствия в папке ./Исходники.\nНе можем удалить эту книгу из ./Размеченные,\
                                       потому что она открыта в Эксель!")

        return ("\n" + ">" + "\n").join(errors_list)

    def make_files(self):

        marked_folder = os.path.join(global_vars.project_folder, r".Размеченные")
        files = [file for file in list(os.walk(os.path.join(global_vars.project_folder, '.Размеченные')))[0][2] if file[0] != "~"]
        columns_info_df = pd.read_excel(os.path.join(global_vars.project_folder,'markup.xlsx'), dtype=str)

        self.res_folder_files_qty = 0
        for file_number, file in enumerate(files, 1):
            with pd.ExcelFile(os.path.join(marked_folder, file)) as xlsx_file:
                sheets = xlsx_file.sheet_names
                    
            for sheet_number, sheet in enumerate(sheets, 1):
                # sleep(0.0001)
                print(file, sheet)
                
                self.mysignal_info_label.emit(
                    f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} '
                    f'Книга {file_number} из {len(files)} лист {sheet_number} из {len(sheets)}. ' 
                    f'Подготавливаем к созданию файл для: "{file}" лист "{sheet}"',
                    'color: blue')
                 
                file_info = columns_info_df[(columns_info_df['_file_'] == file) &
                                            (columns_info_df['_sheet_'] == sheet) &
                                            (columns_info_df['Ошибки маркировки'] == 'ok')]
                
                # print(file_info)
                
                if not file_info.empty:
                    s = int(file_info['_s_'].iloc[0])-1
                    f = int(file_info['_f_'].iloc[0])
                    print(Fore.YELLOW,  os.path.join(marked_folder,file), Fore.RESET)
                    file_df = pd.read_excel(os.path.join(marked_folder, file), sheet_name=sheet, header=None).iloc[:,2:]

        
                    file_df_without_ffill = file_df.iloc[s:f]
                    file_df_without_ffill.columns = file_df.iloc[0]
                    column_names_without_ffill = [column_name for column_name in file_df.iloc[0] if pd.notna(column_name)]
                    file_df_without_ffill = file_df_without_ffill[column_names_without_ffill]

                    # file_df_with_ffill = file_df.fillna(method='ffill')     
                    file_df_with_ffill = file_df.ffill()                       
                    file_df_with_ffill = file_df_with_ffill.iloc[s:f]
                    file_df_with_ffill.columns = file_df.iloc[1]
                    column_names_with_ffill = [column_name for column_name in file_df.iloc[1] if pd.notna(column_name)]
                    file_df_with_ffill = file_df_with_ffill[column_names_with_ffill]           

                    file_df = pd.concat([file_df_without_ffill, file_df_with_ffill], axis=1)


                    res_file_name = os.path.join(global_vars.project_folder, '.Результат', f"{file[3:-4]}_{sheet}.xlsx")
                    file_df.to_excel(res_file_name, index=None)


      
                    # задаём ширину столбцов по размеру заголовка
                    wb = load_workbook(res_file_name)
                    ws = wb.active

                    self.mysignal_info_label.emit(
                        f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                        f"Подгоняем ширину столбцов под длины заголовков",
                        'color: blue')

                    for n, column in enumerate(list(file_df .columns), 1):
                        ws.column_dimensions[get_column_letter(n)].width = len(str(column))*1.1 + 5

                    self.mysignal_info_label.emit(
                        f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                        f"Замораживаем строку заголовков",
                        'color: blue')
                    ws.auto_filter.ref = ws.dimensions    
                    
                    ws.freeze_panes = ws.cell(column=1, row=2)

                    self.mysignal_info_label.emit(
                        f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                        f"Задаём цвет строки заголовков",
                        'color: blue')
                    
                    for col, column in enumerate(list(file_df .columns), start=1):
                        cell = ws.cell(column=col, row = 1)
                        # cell.fill = styles.PatternFill(start_color='FFFFC7CE', fill_type='solid')
                        cell.fill = styles.PatternFill(start_color='E9FDD9', fill_type='solid')
                        cell.font = styles.Font(color='974706', bold=True)
                        cell.alignment = styles.Alignment(wrap_text=True,
                                                        vertical='top',
                                                        horizontal='center') 
                        #cell.style.alignment.wrap_text=True

                    self.mysignal_info_label.emit(
                        f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} '
                        f'Книга {file_number} из {len(files)} лист {sheet_number} из {len(sheets)}. '
                        f'Сохраняем в файл "{file[3:-4]}_{sheet}.xlsx"',
                        'color: blue')    
                    wb.save(res_file_name)
                    self.res_folder_files_qty += 1

        os.startfile(os.path.join(global_vars.project_folder,'.Результат'))                


    def run(self): 


        self.mysignal_info_label.emit(
            f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. "
            f"Проверяем не менялись ли файлы в папке .Исходники",
            'color: blue'
        )
        self.is_src_files_modifyed = check_files_modified('.Исходники')

        self.mysignal_info_label.emit(
            f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. "
            f"Проверяем не менялись ли файлы в папке .Размеченные",
            'color: blue'
        )

        self.is_md_files_modifyed = check_files_modified('.Размеченные')
        print('self.is_md_files_modifyed', self.is_md_files_modifyed)


        # проверяем, чтобы если существует папка .Результат, чтобы не было открытых файлов на рабочем столе
        result_folder = os.path.join(global_vars.project_folder, '.Результат')
        if os.path.exists(result_folder):

            files_list = list(os.walk(result_folder))[0][2]
            files_list = [file for file in files_list if file[:2] != '~$']
            
            if files_list:
            
                files_list.sort()

                self.res_folder_files_opened = []
            
                for file_number, file in enumerate(files_list):
                    self.mysignal_info_label.emit(
                        f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. {file_number} из {len(files_list)}. "
                        f"Создание файлов. Проверяем не открыт ли на рабочем столе: {file}.",
                        'color: blue')
                    
                    #sleep(0.01)
                    if check_excel_file_is_open(file):
                        self.res_folder_files_opened.append(file)
                    else:
                        os.remove(os.path.join(result_folder, file))

                if self.res_folder_files_opened:
                    self.warning_message = (
                        f"Некоторые файлы из папки .Результат открыты на рабочем столе!\n"
                        f"{'\n'.join(self.res_folder_files_opened)}"
                        )

                                  
                    #return
                
        else:

            os.mkdir(result_folder)




        if self.is_src_files_modifyed:

            self.error_message = ('В папку .Исходники были добавлены новые файлы или\n'
                                  'некоторые файлы в ней были пересохранены или удалены.\n'
                                  'Нажмите кнопку "Просмотерь разметку"!')
            #return
        if self.is_md_files_modifyed:

            self.error_message = ('Файлы в папке .Размеченные были изменены.\n'
                                  'Нажмите кнопку "Просмотерь разметку"!!')  
            #return


        self.error_message = ""
        self.warning_message = ""        
        self.error_message = self.clean_folder_marked()

        if not self.error_message:  
            self.make_files()

        self.on_finished()


    def on_clicked(self):     
        self.start() # Запускаем поток  
    

    def on_finished(self): # Вызывается при завершении потока

        if self.error_message:        
            self.mysignal_info_label.emit(
                self.error_message.replace('\n',' '),
                'color: red')
            
            QtWidgets.QMessageBox.critical(None,
                                           self.message_title,
                                           self.error_message,
                                           buttons=QtWidgets.QMessageBox.StandardButton.Ok)
            return
        
        elif self.warning_message:
            pop_up_files(self.message_title, self.warning_message, self.res_folder_files_opened, folder = '.Результат')
            
            self.mysignal_info_label.emit(
                self.warning_message.replace('\n',' '),
                'color: red')
           
        else:

            self.mysignal_info_label.emit(
                f'В папке .Результат создано { self.res_folder_files_qty } файлов.',
                'color: green')

        
        global_vars.ui.pushButtonConcat.setEnabled(True)
        global_vars.ui.pushButtonMakeFiles.setEnabled(True)