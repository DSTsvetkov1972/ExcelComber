from PySide6 import QtWidgets, QtCore
from colorama import Fore
from datetime import datetime
import global_vars 
import os
import pandas as pd
from my_threads.functions import check_files_modified
from openpyxl import load_workbook, styles
from my_threads.functions import all_control_elements_off, all_control_elements_on

class HeadersFillerThread(QtCore.QThread):
 
    mysignal = QtCore.Signal(str)


    def __init__ (self, parent=None):
        QtCore.QThread.__init__(self, parent) 

           
    def check_md_files_available(self):
        """
        Перед началом обработки проверяем чтобы не было 
        открытых размеченных файлов
        """

        self.mysignal.emit(f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                            f"проверяем, чтобы не было открытых файлов из папки .Размеченные")

        md_folder = os.path.join(global_vars.project_folder,'.Размеченные')
        md_files = list(os.walk(md_folder))[0][2]
        
        opened_md_files = [md_file[2:] for md_file in md_files if md_file[0]=='~']


        self.err_list = []
        if opened_md_files:
            
            for file in opened_md_files:
                self.err_list.append(f'{file}, Файл из папки .Размеченные открыт на рабочем столе. Его нужно закрыть!')
                os.startfile(os.path.join(md_folder, file))

            if self.err_list:
                self.error_message = (
                    "Некоторые файлы из папки .Размеченные,\n"
                    "открыты на рабочем столе.")
                
            return False
        return True
            




    def headers_filler(self):
        md_folder = os.path.join(global_vars.project_folder,'.Размеченные')
        md_files = list(os.walk(md_folder))[0][2]

        for file_number, file_name in enumerate(md_files, 1):
            need_to_save = False
            file = os.path.join(md_folder, file_name)

            wb = load_workbook(file)

            for sheet_number, sheet in enumerate(wb.sheetnames, 1):
                self.mysignal.emit(f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                    f"Книга {file_number} из {len(md_files)} лист {sheet_number} из {len(wb.sheetnames)}. Заполняем заголовки для {file_name} из листа {sheet}") 
                df = pd.read_excel(file, sheet_name=sheet, header=None)


                if df.empty:
                    continue

                header_df = df.iloc[:2]

                header_df = header_df.fillna("")

                if not (header_df.applymap(lambda x: isinstance(x, str) and len(x) == 0)).all().all():
                    print('Заголовок уже есть!')
                    continue

                df_0 = df.iloc[0].dropna()

                if df_0.empty:
                    header_df = df[df[0]=='h']

                    if not header_df.empty:
                        need_to_save = True
                        header_cells = header_df.iloc[0][2:]
                    else:
                        continue


                    ws = wb[sheet]

                    col = 3
                    for header_cell in header_cells:
                        ws.cell(row=1, column=col, value=header_cell)
                        ws.cell(row=1, column=col).alignment = styles.Alignment(wrap_text=False, horizontal="center", vertical="center")
                        col += 1

            if need_to_save:
                self.mysignal.emit(f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                f"Книга {file_number} из {len(md_files)}. Сохраняем файл {file_name} после заполнения заголовка.") 
                wb.save(file)      
            else:
                wb.close()





    def on_signal(self,mysignal):          
        global_vars.ui.info_label.setText(mysignal)


    def run(self): 
        self.message_title = "Заполняем заголовки"
        self.error_message = ""
        self.warning_message = ""

        if not self.check_md_files_available():
            return 
        else:
            self.headers_filler()                


    def on_clicked(self):
        self.start() # Запускаем поток  
     


    def on_started(self): # Вызывается при запуске потока
        all_control_elements_off()
        # global_vars.ui.pushButtonChooseProjectFolder.setEnabled(False)   
        global_vars.ui.info_label.setStyleSheet('color: blue')        


    def on_finished(self): # Вызывается при завершении потока
        global_vars.ui.pushButtonChooseProjectFolder.setEnabled(True)
        all_control_elements_on()


        if self.error_message:
            global_vars.ui.info_label.setStyleSheet('color: red')             
            global_vars.ui.info_label.setText(f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                                              f"{self.error_message.replace('\n',' ')}")
            QtWidgets.QMessageBox.critical(None,
                                           self.message_title,
                                           self.error_message,
                                           buttons=QtWidgets.QMessageBox.StandardButton.Ok)
            #refresh_files_info('.Исходники')        
            #refresh_files_info('.Размеченные')             

        elif self.warning_message:
            global_vars.ui.info_label.setStyleSheet('color: red')             
            global_vars.ui.info_label.setText(f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                                              f"{self.warning_message.replace('\n',' ')}")
            
            QtWidgets.QMessageBox.warning(None,
                                           self.message_title,
                                           self.warning_message,
                                           buttons=QtWidgets.QMessageBox.StandardButton.Ok)             
        else:
            global_vars.ui.info_label.setStyleSheet('color: green')             
            global_vars.ui.info_label.setText(
                f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                f"Заголовки заполнены.")
