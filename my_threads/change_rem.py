from PySide6 import QtWidgets, QtCore
from colorama import Fore
from openpyxl import load_workbook
import os
from PySide6 import QtWidgets, QtCore
from colorama import Fore
import global_vars 
import os
import pyperclip
from datetime import datetime
from my_threads.functions import get_files_and_sheets_from_pyperclip, check_files_modified, check_excel_file_is_open, all_control_elements_off, all_control_elements_on, on_finsh_change_thread


class ChangeRemThread(QtCore.QThread):
    def __init__ (self, md_files=False, parent=None):
        QtCore.QThread.__init__(self, parent)
        self.message_title = "Замена примечаний на выбранных листах."
        self.md_files = md_files

    mysignal_info_label = QtCore.Signal(str, str)
    mysignal_finished = QtCore.Signal(str, str, str, str, list)


    def run(self):
        self.error_message = "Не удалось найти ни одного комментария для изменения!"
        self.warning_message = ""
        self.info_message = ""
        self.md_files_opened = []
        changed_qty = 0
        

        self.is_src_files_modifyed = check_files_modified('.Исходники')


        if self.is_src_files_modifyed:

            self.error_message = ('В папку .Исходники были добавлены новые файлы или\n'
                                  'некоторые файлы в ней были пересохранены или удалены.\n'
                                  'Нажмите кнопку "Просмотерь разметку"!')

            # global_vars.ui.info_label.setText(self.error_message)
            return
        
        
        if global_vars.ui.lineEditOldRem.text() == global_vars.ui.lineEditNewRem.text():
            self.error_message = "Старый комментарий такой же как новый!"
            return
    


        files_sheets_list = get_files_and_sheets_from_pyperclip()
        files_list = list({files_sheets[0] for files_sheets in files_sheets_list})
        files_list.sort()        

        self.md_files_opened = [file for file in files_list if check_excel_file_is_open(file)]

        self.md_files_opened = []
        
        for file_number, file in enumerate(files_list):
            self.mysignal_info_label.emit(
                f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. {file_number} из {len(files_list)}. "
                f"Изменение комментариев. Проверяем не открыт ли на рабочем столе: {file}.",
                'color: blue')

            if check_excel_file_is_open(file):
                self.md_files_opened.append(file)
                

        if self.md_files_opened:
            self.warning_message = (
                f"Некоторые размеченные файлы открыты на рабочем столе!\n"
                f"{'\n'.join(self.md_files_opened)}"
                )

            #for md_file in md_files_opened:
            #    os.startfile(os.path.join(global_vars.project_folder, '.Размеченные', md_file))
            return


        file_preceding = ""
        need_to_save = False


        for file_sheet_number, file_sheet_list in enumerate(files_sheets_list, 1):
            print(Fore.MAGENTA, file_sheet_list,  Fore.RESET)
            file = file_sheet_list[0]

            if file != file_preceding:

                if need_to_save:
                    self.mysignal_info_label.emit(
                        f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. { files_list.index(file)+1 } из {len(files_list)}. '
                        f'Сохраняем с измененными комментариями кнгигу "{file}"',
                        'color: blue')
                    # print(Fore.GREEN, 'Мы тут', Fore.RESET)
                    wb.save(os.path.join(global_vars.project_folder, '.Размеченные', file_preceding))

                file_preceding = file
                need_to_save = False

                self.mysignal_info_label.emit(
                    f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. { files_list.index(file)+1 } из {len(files_list)}. '
                    f'Загружаем для изменения комментариев кнгигу "{file}"',
                    'color: blue')
                
                wb = load_workbook(os.path.join(global_vars.project_folder, '.Размеченные', file))
                wb = load_workbook(os.path.join(global_vars.project_folder, '.Размеченные', file))

            
            sheet_name = file_sheet_list[1]


            self.mysignal_info_label.emit(
                f"{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")} "
                f"{file_sheet_number} из {len(files_sheets_list)}. Сканируем примечания в листе { sheet_name } в книге {file}.",
                'color: blue')

            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                self.info_message = f"Изменено примечаний: {changed_qty}."
                continue
                
            rem_in_sheet = '' if not ws['A1'].value else str(ws['A1'].value)
            print(
                file,
                sheet_name,
                Fore.GREEN, str(ws['A1'].value), Fore.RESET,
                Fore.RED, rem_in_sheet, Fore.RESET,
                global_vars.ui.lineEditOldRem.text(),
                global_vars.ui.lineEditNewRem.text(),
                rem_in_sheet == global_vars.ui.lineEditOldRem.text())

            if rem_in_sheet == global_vars.ui.lineEditOldRem.text():
                ws['A1'].value = global_vars.ui.lineEditNewRem.text()
                need_to_save = True
                changed_qty += 1
                self.error_message = ""


        if need_to_save:
            self.mysignal_info_label.emit(
                f'{datetime.strftime(datetime.now(), "%Y-%m-%d %H:%M:%S")}. { files_list.index(file)+1 } из {len(files_list)}. '
                f'Сохраняем кнгигу "{file}"',
                'color: blue')
            wb.save(os.path.join(global_vars.project_folder, '.Размеченные', file_preceding))
        else:
            wb.close()

        self.info_message = f"Изменено примечаний: {changed_qty}."            

        self.mysignal_finished.emit(
            self.message_title,
            self.error_message,
            self.warning_message,
            self.info_message,
            self.md_files_opened
        )



    def on_clicked(self):
        self.old_rem = global_vars.ui.lineEditOldRem.text()
        self.new_rem = global_vars.ui.lineEditNewRem.text()
        self.start() # Запускаем поток  
     


#    def on_finished(self): # Вызывается при завершении потока
#        on_finsh_change_thread(self.message_title, self.error_message, self.warning_message, self.info_message, self.md_files_opened)
#        all_control_elements_on()

